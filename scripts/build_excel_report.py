"""Builds the Excel results workbook from the CSV tables in reports/tables.

    python scripts/build_excel_report.py

Run scripts/run_notebook.py first so the tables are current.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
TABLES = ROOT / "reports" / "tables"
FIGURES = ROOT / "reports" / "figures"

CONFIG = {'file': 'RFM_Segmentation_Results.xlsx',
 'title': 'RFM Customer Segmentation',
 'subtitle': 'Whom to reward, develop and win back',
 'findings': ['Promising (35.4% of customers) and Cannot Lose Them (16.2%) hold 80.5% of revenue.',
              '88.7% of customers placed one or two orders, so segments are driven mainly by recency and '
              'spend.',
              'Cannot Lose Them customers have not ordered for 979 days on average.',
              'Champions and Loyal customers are 0.6% of customers but spend about three times more per '
              'head.',
              'Need Attention: 2.2% of customers, 7.0% of revenue, $1,007 per customer.'],
 'sheets': {'Segments': ('segment_summary.csv', 'customers, revenue and behaviour by segment'),
            'Customer scores': ('customer_rfm_scores.csv', 'R, F, M scores and segment for every customer')},
 'figure': '02_segment_treemaps.png',
 'source': 'Synthetic educational e-commerce data, 51,290 order lines, 2020-2023'}

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16, color="1F4E79")


def number_format(column, series):
    name = column.lower()
    if not pd.api.types.is_numeric_dtype(series):
        return None
    if name.endswith("pct") or "_pct" in name or name.endswith("rate"):
        return '0.0"%"'
    if any(key in name for key in ("items", "orders", "sessions", "customers", "titles", "count")):
        return "#,##0"
    if any(key in name for key in ("usd", "revenue", "profit", "monetary", "refund", "margin", "aov", "per_session",
                                   "per_customer")):
        return '"$"#,##0.00' if series.abs().max() < 100 else '"$"#,##0'
    if pd.api.types.is_integer_dtype(series):
        return "#,##0"
    return "#,##0.00"


def add_table_sheet(wb, title, frame):
    ws = wb.create_sheet(title[:31])
    ws.append(list(frame.columns))
    for row in frame.itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.item() if hasattr(v, "item") else v) for v in row])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for idx, column in enumerate(frame.columns, start=1):
        fmt = number_format(column, frame[column])
        letter = get_column_letter(idx)
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        width = max(len(str(column)), *(len(str(v)) for v in frame[column].head(200))) if len(frame) else len(column)
        ws.column_dimensions[letter].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = CONFIG["title"]
    ws["A1"].font = TITLE_FONT
    ws["A2"] = CONFIG["subtitle"]
    ws["A2"].font = Font(italic=True, color="555555")
    row = 4
    ws.cell(row, 1, "Key findings").font = Font(bold=True, size=12)
    for finding in CONFIG["findings"]:
        row += 1
        ws.cell(row, 1, f"• {finding}").alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
    row += 2
    ws.cell(row, 1, "Sheets in this workbook").font = Font(bold=True, size=12)
    for sheet, (_, description) in CONFIG["sheets"].items():
        row += 1
        ws.cell(row, 1, f"{sheet}: {description}")
    row += 2
    ws.cell(row, 1, "Source: " + CONFIG["source"]).font = Font(color="555555")
    ws.column_dimensions["A"].width = 120
    if CONFIG.get("figure") and (FIGURES / CONFIG["figure"]).exists():
        image = Image(str(FIGURES / CONFIG["figure"]))
        scale = 620 / image.width
        image.width, image.height = image.width * scale, image.height * scale
        ws.add_image(image, f"A{row + 2}")

    for sheet, (csv_name, _) in CONFIG["sheets"].items():
        add_table_sheet(wb, sheet, pd.read_csv(TABLES / csv_name))

    out = ROOT / "reports" / CONFIG["file"]
    wb.save(out)
    print(f"Wrote {out.relative_to(ROOT)} with {len(wb.sheetnames)} sheets")


if __name__ == "__main__":
    main()
